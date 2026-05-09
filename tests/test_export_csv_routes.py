from pathlib import Path
import io
import sys

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.append(str(Path(__file__).resolve().parents[1] / "src"))

import api.routes as routes
from config import AppConfig


class StubDbcSignal:
    def __init__(self, name: str, unit: str | None):
        self.name = name
        self.unit = unit


class StubDbcMessage:
    def __init__(self, name: str, signals):
        self.name = name
        self.signals = signals


class StubDbc:
    def __init__(self):
        self.messages = {
            "VehicleState": StubDbcMessage(
                "VehicleState",
                [
                    StubDbcSignal("speed", "mph"),
                    StubDbcSignal("gear", None),
                ],
            )
        }

    def get_message_by_name(self, name: str):
        return self.messages[name]


class StubDbcManager:
    def __init__(self, db=None):
        self.db = db

    def get_active_dbc(self):
        return self.db


class StubLogFileDataSource:
    instances = []

    def __init__(self, log_file_path: str, playback_speed: float = 1.0, db=None):
        self.log_file_path = log_file_path
        self.playback_speed = playback_speed
        self.db = db
        self.connected = False
        self.disconnected = False
        StubLogFileDataSource.instances.append(self)

    async def connect(self):
        self.connected = True

    async def disconnect(self):
        self.disconnected = True

    async def stream(self):
        yield {
            "timestamp": 1.25,
            "message_name": "VehicleState",
            "decoded": {
                "speed": 12.5,
                "gear": 3,
            },
        }
        yield {
            "timestamp": 2.5,
            "message_name": "VehicleState",
            "decoded": {
                "speed": 13.0,
                "abcdefghijklmnopqrstuvwxyz123456789": 8.8,
                "abcdefghijklmnopqrstuvwxyz123456780": 9.9,
            },
        }
        yield {
            "timestamp": 3.0,
        }


def build_client(tmp_path: Path, dbc=None, log_file_exists: bool = True):
    app = FastAPI()
    app.include_router(routes.router)
    app.state.config = AppConfig()
    log_file = tmp_path / "sample.blf"
    if log_file_exists:
        log_file.write_bytes(b"blf")
    app.state.config.pipeline.log_file = str(log_file)
    app.state.dbc_manager = StubDbcManager(dbc)
    return TestClient(app), log_file


def test_export_csv_requires_configured_log_file(tmp_path: Path):
    client, _ = build_client(tmp_path, dbc=StubDbc(), log_file_exists=False)

    response = client.get("/api/export_csv")

    assert response.status_code == 400
    assert response.json() == {"detail": "No configured log file was found."}


def test_export_csv_requires_active_dbc(tmp_path: Path):
    client, _ = build_client(tmp_path, dbc=None)

    response = client.get("/api/export_csv")

    assert response.status_code == 400
    assert response.json() == {"detail": "No active DBC is loaded."}


def test_export_csv_rejects_non_us_timezones(tmp_path: Path):
    client, _ = build_client(tmp_path, dbc=StubDbc())

    response = client.get("/api/export_csv", params={"timezone_name": "Europe/London"})

    assert response.status_code == 400
    assert response.json() == {"detail": "Timezone must be one of the supported United States options."}


def test_export_csv_streams_bom_csv_and_reuses_logfile_source(monkeypatch, tmp_path: Path):
    StubLogFileDataSource.instances.clear()
    monkeypatch.setattr(routes, "LogFileDataSource", StubLogFileDataSource)
    client, log_file = build_client(tmp_path, dbc=StubDbc())

    response = client.get("/api/export_csv", params={"timezone_name": "America/Chicago"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv; charset=utf-8")
    assert "attachment; filename=\"wolftrack-export-" in response.headers["content-disposition"]
    assert response.text.startswith("\ufefftimestamp,signal_name,value,unit\n")
    assert "1969-12-31T18:00:01.250000-06:00,speed,12.5,mph\n" in response.text
    assert "1969-12-31T18:00:01.250000-06:00,gear,3,\n" in response.text
    assert "1969-12-31T18:00:02.500000-06:00,speed,13.0,mph\n" in response.text

    assert len(StubLogFileDataSource.instances) == 1
    instance = StubLogFileDataSource.instances[0]
    assert instance.log_file_path == str(log_file)
    assert instance.playback_speed == 0
    assert instance.connected is True
    assert instance.disconnected is True


def test_export_xlsx_rejects_non_us_timezones(tmp_path: Path):
    client, _ = build_client(tmp_path, dbc=StubDbc())

    response = client.get("/api/export_xlsx", params={"timezone_name": "Europe/London"})

    assert response.status_code == 400
    assert response.json() == {"detail": "Timezone must be one of the supported United States options."}


def test_export_xlsx_returns_grouped_workbook(monkeypatch, tmp_path: Path):
    StubLogFileDataSource.instances.clear()
    monkeypatch.setattr(routes, "LogFileDataSource", StubLogFileDataSource)
    client, log_file = build_client(tmp_path, dbc=StubDbc())

    response = client.get("/api/export_xlsx", params={"timezone_name": "America/Chicago"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=\"wolftrack-export-" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith(".xlsx\"")

    workbook = load_workbook(io.BytesIO(response.content))
    assert "speed" in workbook.sheetnames
    assert "gear" in workbook.sheetnames
    assert "abcdefghijklmnopqrstuvwxyz12345" in workbook.sheetnames
    assert "abcdefghijklmnopqrstuvwxyz123_1" in workbook.sheetnames

    speed_sheet = workbook["speed"]
    assert [cell.value for cell in speed_sheet[1]] == ["timestamp", "value", "unit"]
    assert speed_sheet["A1"].font.bold is True
    assert speed_sheet["B1"].font.bold is True
    assert speed_sheet["C1"].font.bold is True
    assert speed_sheet["A2"].value == "1969-12-31T18:00:01.250000-06:00"
    assert speed_sheet["B2"].value == 12.5
    assert speed_sheet["C2"].value == "mph"
    assert speed_sheet["A3"].value == "1969-12-31T18:00:02.500000-06:00"
    assert speed_sheet["B3"].value == 13
    assert speed_sheet["C3"].value == "mph"
    assert speed_sheet.column_dimensions["A"].width > len("timestamp")

    gear_sheet = workbook["gear"]
    assert [cell.value for cell in gear_sheet[1]] == ["timestamp", "value", "unit"]
    assert gear_sheet["B2"].value == 3
    assert gear_sheet["C2"].value in ("", None)

    long_sheet = workbook["abcdefghijklmnopqrstuvwxyz12345"]
    assert long_sheet["B2"].value == 8.8
    assert long_sheet["C2"].value in ("", None)

    collision_sheet = workbook["abcdefghijklmnopqrstuvwxyz123_1"]
    assert collision_sheet["B2"].value == 9.9
    assert collision_sheet["C2"].value in ("", None)

    assert len(StubLogFileDataSource.instances) == 1
    instance = StubLogFileDataSource.instances[0]
    assert instance.log_file_path == str(log_file)
    assert instance.playback_speed == 0
    assert instance.connected is True
    assert instance.disconnected is True
